from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = 'tokenomics-ai-native-currency-research.docx'
BLUE = '2E74B5'; NAVY = '0B2545'; LIGHT = 'E8EEF5'; GRAY = 'F2F4F7'; MUTED = '5B6573'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); m = tcPr.first_child_found_in('w:tcMar')
    if m is None: m = OxmlElement('w:tcMar'); tcPr.append(m)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        x = m.find(qn('w:'+side))
        if x is None: x=OxmlElement('w:'+side); m.append(x)
        x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa')
def set_cell(cell, text, bold=False, color='000000', size=9.5):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
    r=p.add_run(text); r.bold=bold; r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(cell)
def set_table_width(table, widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr; tblW=tblPr.first_child_found_in('w:tblW'); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); tblPr.append(ind)
    dxa = [round(w * 1440) for w in widths]
    for col, width in zip(table._tbl.tblGrid.gridCol_lst, dxa):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for i,w in enumerate(widths):
            row.cells[i].width=Inches(w)
            tcPr = row.cells[i]._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in('w:tcW')
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'), str(dxa[i])); tcW.set(qn('w:type'), 'dxa')
    trPr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)
def add_bullets(doc, items):
    for t in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.15; p.add_run(t)
def add_heading(doc, text, level=1): doc.add_heading(text, level=level)
def add_para(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.1
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else:p.add_run(text)
    return p
def link_text(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
    r=p.add_run(text); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.75); sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=Inches(.49); sec.footer_distance=Inches(.49)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for n,size,col,before,after in [('Heading 1',16,BLUE,12,6),('Heading 2',13,BLUE,10,5),('Heading 3',11,NAVY,8,4)]:
 s=styles[n]; s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
header=sec.header.paragraphs[0]; header.text='TOKENOMICS RESEARCH BRIEF'; header.runs[0].font.name='Arial'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run('22 August 2026  |  Research draft') .font.size=Pt(8)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(28); p.paragraph_format.space_after=Pt(4); r=p.add_run('TOKENOMICS DECISION BRIEF'); r.font.name='Arial'; r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); r=p.add_run('An AI-native, work-released currency'); r.font.name='Arial'; r.font.size=Pt(27); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph('A research-backed proposal for issuance, treasury, fees and agent payments.'); p.runs[0].font.size=Pt(13); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED); p.paragraph_format.space_after=Pt(18)
t=doc.add_table(rows=1, cols=2); set_table_width(t,[1.35,5.15]); set_cell(t.cell(0,0),'Decision',True,NAVY,10); set_cell(t.cell(0,1),'Use a capped, work-released native utility token; use stablecoins for agent payments.',False,'000000',10); shade(t.cell(0,0),LIGHT); shade(t.cell(0,1),'F7F9FC')
add_heading(doc,'The answer in one page')
add_para(doc,'Your instinct is directionally right: reserve most supply for demonstrated network utility. But do not make supply an open-ended function of workload. That creates an incentive to manufacture workload. Instead, work should unlock a bounded reserve under hard epoch caps and independent verification.')
add_bullets(doc,[
'Small initial treasury, not a large pre-mine: 5% illustrative genesis allocation, mostly time-locked and publicly budgeted.',
'Most issuance is released only after a paid, verified result: the Work Receipt is the economic primitive.',
'Use a dynamic base protocol fee and burn it; do not rely on burn as the primary inflation-control mechanism.',
'Use stablecoins for agent pricing and budgets. Make the native token the assurance, staking and reward asset - not a volatile tollbooth for every API call.'
])
add_heading(doc,'Recommended design: Proof-of-Useful-Settlement',1)
add_para(doc,'A completed job creates a signed Work Receipt: request commitment, stablecoin quote/payment attestation, execution evidence, output hash, buyer acceptance (or expiry), independent quality score and dispute result. Only receipts with an independent buyer, a unique request, task-specific verification, and sufficient bond/reputation may unlock rewards.')
add_para(doc,'Keep prompts, inputs and outputs private; store commitments and selective audit evidence rather than raw AI data. The system rewards paid and verified settlement, not GPU seconds, prompts, token transfers, or self-reported activity.')
add_heading(doc,'Illustrative supply policy',1)
t=doc.add_table(rows=1, cols=3); set_table_width(t,[1.35,1.0,4.15])
for i,x in enumerate(['Pool','Share','Release rule']): set_cell(t.cell(0,i),x,True,NAVY,9.5); shade(t.cell(0,i),LIGHT)
rows=[('Genesis treasury','5%','1% operating runway, 24-month linear release; 4% governed R&D/community treasury, timelocked. No mint key.'),('Verified work','55%','70% slow time release; 30% only as public useful-work baseline is reached. Unreleased supply is not discretionary treasury.'),('Security/verification','25%','Validators, challengers, availability and disputes. Predictable decay; review a tail emission only before it begins.'),('Ecosystem','15%','Grants, interoperability and early liquidity with on-chain vesting and public recipient disclosure.')]
for row in rows:
 c=t.add_row().cells
 for i,x in enumerate(row): set_cell(c[i],x,i==0,'000000',9)
add_para(doc,'Illustration only: use a supply simulator before setting any numbers. The meaningful commitments are the maximum supply, release caps, vesting, disclosure, and absence of privileged minting.')

doc.add_page_break(); add_heading(doc,'How issuance should actually work',1)
add_para(doc,'Set a hard maximum supply (the report uses 1 billion units only as an example). Each epoch releases:')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(10); r=p.add_run('work release = min(epoch cap, reserve remaining, quality-adjusted settlement score x release factor)'); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
add_para(doc,'The epoch cap is the crucial constraint. Cap scores by buyer, worker, verifier group, task family and epoch; discount correlated identities; and delay finality for disputes. Start with mostly time-based release until scoring has survived live adversarial testing.')
add_heading(doc,'Why raw workload fails',2)
add_bullets(doc,[
'A worker can generate its own jobs, recycle prompts, over-compute, or create Sybil buyers.',
'Subjective AI quality cannot be inferred from cost or token count; it needs validation and a challenge path.',
'A work metric becomes the target. Anything measurable must be assumed gameable until the economics make gaming uneconomic.'
])
add_heading(doc,'Controls needed before rewards expand',2)
add_bullets(doc,[
'Independent buyer requirement, buyer deposits for high-value work, and limits on affiliated activity.',
'Staking plus slashing for workers and verifiers; random re-execution and paid challengers.',
'Task-specific proofs where possible: deterministic replay, benchmark suites, TEEs, or zero-knowledge verification. Use subjective validator consensus only where objective checking is impossible.',
'A public receipts schema, fraud taxonomy, audit dashboard and worst-case supply simulator.'
])
add_heading(doc,'Conventional options compared',1)
t=doc.add_table(rows=1, cols=3); set_table_width(t,[1.55,2.45,2.5])
for i,x in enumerate(['Model','Strength','Failure mode / fit']):set_cell(t.cell(0,i),x,True,NAVY,9.5);shade(t.cell(0,i),LIGHT)
for row in [('Fixed halving / hard cap','Simple, credible and easy to model.','Pays even when demand is zero; still needs a security budget.'),('Pure fee-only','No dilution after launch.','Underfunds early security and verification before usage exists.'),('Uncapped work inflation','Strong immediate worker incentive.','Rewards manufactured activity; supply and agent economics become unpredictable.'),('Governance-set inflation','Flexible for changing conditions.','Weak credible neutrality and easy political capture.'),('Capped work release (recommended)','Connects issuance to utility while bounding supply.','Needs a conservative verifier and anti-fraud design.')]:
 c=t.add_row().cells
 for i,x in enumerate(row):set_cell(c[i],x,i==0,'000000',8.7)

doc.add_page_break(); add_heading(doc,'Burns, fees and treasury discipline',1)
add_heading(doc,'Burn a base fee; do not burn worker income',2)
add_para(doc,'Adopt a dynamic base protocol fee for settlement/verification bandwidth and burn it. Charge a separate explicit priority/service fee to compensate block producers and verifiers. This mirrors the valuable separation in EIP-1559: the party selecting transactions does not receive the base fee it can influence.')
add_bullets(doc,[
'Inflation control comes first from the maximum supply, epoch caps and vesting - not from burn forecasts.',
'Burn is demand-dependent: quiet periods may burn almost nothing. Never market a deflation promise.',
'For slashing, pay a portion to the successful challenger and burn a portion. Detection needs an economic reward.'
])
add_heading(doc,'Treasury principles',2)
add_bullets(doc,[
'Publish balances, unlock calendar, wallets, spending mandate, grants and recipient disclosures.',
'Use a timelock, multisignature and narrow emergency powers. An emergency should pause a subsystem, not create new tokens.',
'Keep team and foundation compensation separate from the work reserve. Avoid ambiguous labels such as “ecosystem” that hide discretionary dilution.'
])
add_heading(doc,'AI-friendly currency: two rails',1)
t=doc.add_table(rows=1, cols=3);set_table_width(t,[1.35,2.4,2.75])
for i,x in enumerate(['Rail','Use','Why it matters']):set_cell(t.cell(0,i),x,True,NAVY,9.5);shade(t.cell(0,i),LIGHT)
for row in [('Stable payment rail','USDC-style quotes and x402-compatible HTTP 402 settlement for APIs, data, compute and agent services.','Agents can set real budgets and compare providers without exposure to native-token volatility.'),('Native assurance rail','Stake, validation rewards, challenge bonds, slashing, Work Receipt rewards and constrained governance.','The token has a necessary security role without forcing every user to speculate before buying a service.')]:
 c=t.add_row().cells
 for i,x in enumerate(row):set_cell(c[i],x,i==0,'000000',9)
add_heading(doc,'What is actually unique',2)
add_para(doc,'Not “AI coin.” The differentiator is a portable Work Receipt: a cryptographically signed, privacy-preserving proof that a named agent bought or delivered a verified service at an agreed price. It can feed reputation, unlock a bounded work reserve, support disputes, and allow agents to accumulate a credible service history across marketplaces.')
add_para(doc,'Pair that receipt with portable agent identity, reputation and validation. Keep trust separate from payments: payments settle value; validation establishes whether the value was delivered.')

doc.add_page_break(); add_heading(doc,'Launch roadmap and decision gates',1)
for title,body in [('0. Simulate and specify','Publish issuance, fee/burn and treasury scenarios under low, base and high demand. Release the Work Receipt schema, a threat model and audit plan before token distribution.'),('1. Objective tasks only','Reward deterministic compute, data availability or benchmarkable inference first. Apply low caps across buyers, workers, verifier groups and task classes.'),('2. Challenge market','Before subjective AI tasks, prove random audits, re-execution, bonds, appeals and slashing in production-like adversarial tests.'),('3. Controlled decentralisation','Put supply cap, treasury unlocks and burn logic behind long timelocks and supermajority thresholds. Let task scoring evolve faster, but only within bounded budget envelopes.')]:
 add_heading(doc,title,2);add_para(doc,body)
add_heading(doc,'Decision checklist',1)
add_bullets(doc,[
'Can a related party manufacture paid demand cheaply enough to profit after fees, bonds and expected slashing?',
'Can an honest verifier profitably challenge bad work?',
'Can an agent quote, authorize and reconcile a service purchase in fiat terms without holding the native token?',
'Is the supply schedule calculable from published state without trusting a foundation?',
'Would the system still pay for security and dispute resolution if usage fell sharply for a year?'
])
add_heading(doc,'Legal and communications caution',1)
add_para(doc,'This is not legal or financial advice. In the United States, whether a crypto-asset transaction is an investment contract depends on facts, distribution and promotion. Avoid promises of appreciation, “deflationary” marketing and passive return framing; obtain specialist advice in every target jurisdiction before distribution.')
add_heading(doc,'Research sources',1)
for x in [
'Filecoin Docs, “Crypto-economics” (accessed 22 Aug 2026): https://docs.filecoin.io/basics/what-is-filecoin/crypto-economics',
'Ethereum, EIP-1559: Fee market change for ETH 1.0 chain (2019): https://eips.ethereum.org/EIPS/eip-1559',
'Bittensor Docs, “Emissions” (accessed 22 Aug 2026): https://www.bittensor.com/docs/concepts/emissions',
'Coinbase Developer Documentation, “Welcome to x402” (accessed 22 Aug 2026): https://docs.cdp.coinbase.com/x402/welcome',
'Ethereum, ERC-8004: Trustless Agents (draft, 2025): https://eips.ethereum.org/EIPS/eip-8004',
'ethereum.org, “AI agents” (accessed 22 Aug 2026): https://ethereum.org/ai-agents/',
'U.S. SEC, “Transactions Involving Crypto Assets” (Apr. 2026): https://www.sec.gov/resources-small-businesses/capital-raising-building-blocks/transactions-involving-crypto-assets'
]: link_text(doc,x)
doc.save(OUT)
print(OUT)
